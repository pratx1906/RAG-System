"""
Production-grade document parser.

Improvements over original:
- All regex patterns precompiled at module load (not per-line / per-call)
- Excel: read_only=True for memory-safe streaming of large workbooks
- CSV: df.to_dict('records') replaces slow iterrows()
- _row_to_natural_language: O(1) field lookup via normalised key dict (was O(n*m))
- Cell value length capped to prevent single-field blowouts
- Structured row text capped to prevent runaway row chunks
- Improved error handling and parser fallback hygiene
- Text/Markdown heading regex compiled once at module level
"""

import logging
import re
from pathlib import Path
from typing import Dict

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".txt", ".csv", ".md"}

# ── Per-cell / per-row length guards ─────────────────────────────────────────

_MAX_CELL_CHARS = 300   # individual spreadsheet cell value cap
_MAX_ROW_TEXT   = 900   # combined structured text for one spreadsheet row

# ── Precompiled patterns ──────────────────────────────────────────────────────

# Markdown → plain text
_RE_MD_IMG       = re.compile(r"!\[(.*?)\]\((.*?)\)")
_RE_MD_TABLE_SEP = re.compile(r"^[\|\s\-:]+$")
_RE_MD_HEADING   = re.compile(r"^(#{1,6})\s+(.*)")
_RE_MD_BOLD_I    = re.compile(r"\*{1,3}(.*?)\*{1,3}")
_RE_MD_UNDER_I   = re.compile(r"_{1,3}(.*?)_{1,3}")
_RE_MD_STRIKE    = re.compile(r"~~(.*?)~~")
_RE_MD_CODE      = re.compile(r"`(.+?)`")
_RE_MD_LINK      = re.compile(r"\[([^\]]*)\]\([^\)]*\)")
_RE_MULTI_NL     = re.compile(r"\n{3,}")

# Text file heading conversion (multiline: match start of each line)
_RE_TEXT_HEADING = re.compile(r"^#{1,6}\s+(.*)", re.MULTILINE)

# Page anchor detection
_RE_PAGE_ANCHOR  = re.compile(r"^page\s+", re.IGNORECASE)


# ── Document router ───────────────────────────────────────────────────────────

def parse_document(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    parsers = {
        ".pdf":  _parse_pdf,
        ".docx": _parse_docx,
        ".doc":  _parse_docx,
        ".xlsx": _parse_excel,
        ".xls":  _parse_excel,
        ".pptx": _parse_pptx,
        ".csv":  _parse_csv,
        ".txt":  _parse_text,
        ".md":   _parse_text,
    }
    parser = parsers.get(ext)
    if parser is None:
        raise ValueError(f"Unsupported file type: {ext}")
    return parser(file_path)


# ── PDF (Marker → PyPDF2 fallback) ───────────────────────────────────────────

_marker_converter = None


def _get_marker_converter():
    global _marker_converter
    if _marker_converter is not None:
        return _marker_converter
    try:
        from marker.converters.pdf import PdfConverter
        from marker.models import create_model_dict
        _marker_converter = PdfConverter(artifact_dict=create_model_dict())
    except ImportError as exc:
        raise ImportError(
            "marker-pdf is not installed. Install with: pip install marker-pdf"
        ) from exc
    return _marker_converter


def _parse_pdf(path: str) -> str:
    try:
        converter = _get_marker_converter()
        rendered  = converter(path)
        if hasattr(rendered, "markdown"):
            markdown_text = rendered.markdown
        else:
            from marker.output import text_from_rendered
            markdown_text, _, _ = text_from_rendered(rendered)
        return _markdown_to_plain_text(markdown_text)
    except ImportError:
        return _parse_pdf_pypdf2(path)


def _parse_pdf_pypdf2(path: str) -> str:
    import PyPDF2
    parts = []
    with open(path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                parts.append(f"\n[PAGE {i + 1}]\n{text.strip()}")
    return "\n".join(parts)


# ── Markdown → plain text ─────────────────────────────────────────────────────

def _markdown_to_plain_text(markdown: str) -> str:
    """
    Convert Marker-produced markdown to structured plain text.
    All regex patterns are precompiled at module level — not recompiled per line.
    """
    lines        = markdown.split("\n")
    result       = []
    in_table     = False
    page_counter = 1

    for line in lines:
        stripped = line.strip()

        # Page anchors
        if _RE_PAGE_ANCHOR.match(stripped):
            result.append(f"\n[PAGE {page_counter}]\n")
            page_counter += 1
            continue

        # Figures / images
        img_match = _RE_MD_IMG.match(stripped)
        if img_match:
            caption = img_match.group(1).strip()
            result.append(f"\n[FIGURE] {caption}\n" if caption else "\n[FIGURE]\n")
            continue

        # Table separator row
        if "|" in stripped and _RE_MD_TABLE_SEP.fullmatch(stripped):
            in_table = True
            continue

        # Headings
        heading = _RE_MD_HEADING.match(line)
        if heading:
            result.append(f"\n[SECTION] {heading.group(2).strip()}\n")
            continue

        # Strip inline markdown
        line = _RE_MD_BOLD_I.sub(r"\1", line)
        line = _RE_MD_UNDER_I.sub(r"\1", line)
        line = _RE_MD_STRIKE.sub(r"\1", line)
        line = _RE_MD_CODE.sub(r"\1", line)
        line = _RE_MD_LINK.sub(r"\1", line)

        # Table rows
        if "|" in line:
            parts = [p.strip() for p in line.split("|") if p.strip()]
            if parts:
                if not in_table:
                    result.append("\n[TABLE]")
                    in_table = True
                line = " | ".join(parts)
        else:
            in_table = False

        result.append(line)

    text = "\n".join(result)
    text = _RE_MULTI_NL.sub("\n\n", text)
    return text.strip()


# ── DOCX ──────────────────────────────────────────────────────────────────────

def _parse_docx(path: str) -> str:
    from docx import Document as DocxDocument
    doc   = DocxDocument(path)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        parts.append("[TABLE]")
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


# ── Shared: row → natural language ───────────────────────────────────────────

def _row_to_natural_language(fields: Dict[str, str]) -> str:
    """
    Convert a key→value dict into a readable sentence.
    Uses an O(1) normalised-key lookup dict instead of the original O(n×m) scan.
    """
    norm = {k.lower().strip(): v for k, v in fields.items()}

    def get(*keys: str) -> str:
        for k in keys:
            v = norm.get(k, "")
            if v:
                return v
        return ""

    project  = get("project", "name", "title", "task")
    people   = get("people involved", "people", "team", "members",
                   "assigned to", "owner", "responsible")
    status   = get("status", "state", "progress", "stage")
    category = get("category", "type", "domain", "area", "department")
    details  = get("details", "description", "notes", "summary", "objective")
    date     = get("date", "deadline", "due date", "due", "milestone", "timeline")

    parts = []
    if people and project:
        verb          = "is" if "," not in people else "are"
        status_part   = f", Status: {status}"    if status   else ""
        category_part = f", Category: {category}" if category else ""
        parts.append(
            f"{people} {verb} working on {project} (Status: {status}{category_part})."
            if status
            else f"{people} {verb} working on {project}{category_part}."
        )
    elif project:
        parts.append(f"Project: {project}.")

    if details:
        parts.append(f"Details: {details}.")
    if date:
        parts.append(f"Date: {date}.")
    if not parts:
        parts = [f"The {k} is {v}." for k, v in fields.items() if v]

    return " ".join(parts)


def _cap_cell(value: str) -> str:
    """Truncate a single cell value so no one field dominates the chunk text."""
    return value[:_MAX_CELL_CHARS] if len(value) > _MAX_CELL_CHARS else value


# ── Excel ─────────────────────────────────────────────────────────────────────

def _parse_excel(path: str) -> str:
    """
    Parse Excel workbooks.
    read_only=True streams the file rather than loading it fully into memory.
    Cell values and combined row text are capped to prevent oversized chunks.
    """
    from openpyxl import load_workbook

    wb    = load_workbook(path, data_only=True, read_only=True)
    parts = []

    _PRIORITY = ["people involved", "project", "status", "category"]

    for sheet in wb.sheetnames:
        ws       = wb[sheet]
        row_iter = ws.iter_rows(values_only=True)

        try:
            header_row = next(row_iter)
        except StopIteration:
            continue

        headers = [str(h).strip() if h is not None else "" for h in header_row]

        for raw_row in row_iter:
            cells = [_cap_cell(str(v).strip()) if v is not None else "" for v in raw_row]

            if not any(cells):
                continue

            field_dict = {h: v for h, v in zip(headers, cells) if h and v}
            if not field_dict:
                continue

            # Priority-field reordering
            reordered: Dict[str, str] = {}
            low_fields = {k.lower().strip(): k for k in field_dict}
            for pf in _PRIORITY:
                orig_key = low_fields.get(pf)
                if orig_key and orig_key not in reordered:
                    reordered[orig_key] = field_dict[orig_key]
            for k, v in field_dict.items():
                if k not in reordered:
                    reordered[k] = v

            structured = "[Sheet: {}] {} | ".format(
                sheet,
                " | ".join(f"{h}: {v}" for h, v in reordered.items()),
            )
            if len(structured) > _MAX_ROW_TEXT:
                structured = structured[:_MAX_ROW_TEXT].rsplit(" | ", 1)[0] + " | "

            nl = _row_to_natural_language(reordered)
            parts.append(f"{structured}\n{nl}")

    wb.close()
    return "\n---\n".join(parts)


# ── PPTX ──────────────────────────────────────────────────────────────────────

def _parse_pptx(path: str) -> str:
    from pptx import Presentation
    prs  = Presentation(path)
    text = []

    for i, slide in enumerate(prs.slides):
        text.append(f"\n[SLIDE {i + 1}]\n")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text.append(shape.text.strip())

            if shape.has_table:
                text.append("[TABLE]")
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        text.append(row_text)

            if shape.has_chart:
                try:
                    title = shape.chart.chart_title.text_frame.text
                except Exception:
                    title = ""
                text.append(f"[CHART] {title}" if title else "[CHART]")

    return "\n".join(text)


# ── CSV ───────────────────────────────────────────────────────────────────────

def _parse_csv(path: str) -> str:
    """
    Parse CSV files.
    df.to_dict('records') is ~50-100× faster than iterrows() for large files.
    Cell values and combined row text are capped.
    """
    import pandas as pd

    try:
        df = pd.read_csv(path, dtype=str)
    except Exception as exc:
        log.warning("[parser] CSV read failed, retrying with latin-1: %s", exc)
        df = pd.read_csv(path, dtype=str, encoding="latin-1", on_bad_lines="skip")

    parts = []
    for row_num, record in enumerate(df.to_dict("records"), start=1):
        field_dict = {
            col: _cap_cell(str(val).strip())
            for col, val in record.items()
            if val is not None and str(val).strip() not in ("", "nan", "NaN", "None")
        }
        if not field_dict:
            continue

        structured = " | ".join(f"{k}: {v}" for k, v in field_dict.items()) + " | "
        if len(structured) > _MAX_ROW_TEXT:
            structured = structured[:_MAX_ROW_TEXT].rsplit(" | ", 1)[0] + " | "

        nl = _row_to_natural_language(field_dict)
        parts.append(f"[ROW {row_num}]\n{structured}\n{nl}")

    return "\n---\n".join(parts)


# ── Plain text / Markdown ─────────────────────────────────────────────────────

def _parse_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    content = _RE_TEXT_HEADING.sub(r"\n[SECTION] \1\n", content)
    return content
